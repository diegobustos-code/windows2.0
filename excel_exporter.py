# -*- coding: utf-8 -*-
"""
excel_exporter.py
==================
Exporta los resultados del cotejo a archivos .xlsx.

Todos los reportes comparten el mismo formato:
  - Una fila por socio, EN EL MISMO ORDEN del listado oficial.
  - Cada fila lleva el color de su sede (ver sede_colors.py).
  - Celdas de asistencia, todas en negrita:
        P = letra verde, sin relleno propio (se ve el color de la sede)
        A = letra roja sobre relleno rojo suave
        J = letra naranja sobre relleno amarillo
  - Las listas para revisar manualmente van en OTRAS PESTAÑAS DEL MISMO
    LIBRO: "Ambiguos", "No Encontrados" y "Cotejos aproximados".
"""

from copy import copy
from typing import Dict, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from sede_colors import build_sede_fills, fill_for

COLUMNS = ["Nombre", "Apellido", "Sede", "Duración (minutos)", "Asistencia"]

# Formato de las celdas de asistencia (todo en negrita):
#   P = Presente  -> letra verde, SIN relleno propio (se conserva el relleno original)
#   A = Ausente   -> letra roja fuerte sobre relleno rojo suave
#   J = Justificada -> letra naranja sobre relleno amarillo
_FONT_PRESENTE = Font(color="FF008000", bold=True)
_FILL_AUSENTE = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
_FONT_AUSENTE = Font(color="FFC00000", bold=True)
_FILL_JUSTIFICADA = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
_FONT_JUSTIFICADA = Font(color="FFE36C09", bold=True)

SHEET_AMBIGUOS = "Ambiguos"
SHEET_NO_ENCONTRADOS = "No Encontrados"
SHEET_APROXIMADOS = "Cotejos aproximados"


def apply_status_style(cell, base_fill=None) -> bool:
    """
    Aplica el formato de asistencia a una celda cuyo valor sea P, A o J
    (sin distinguir mayúsculas). Devuelve True si la celda tenía uno de
    esos valores y fue formateada; False si no (no se toca nada).

    `base_fill`: relleno "original" a conservar en las P (por ejemplo el
    color de la fila de su sede). Si es None, la P queda sin relleno.
    """
    value = str(cell.value).strip().upper() if cell.value is not None else ""
    if value not in ("P", "A", "J"):
        return False
    cell.alignment = Alignment(horizontal="center")
    if value == "P":
        cell.font = _FONT_PRESENTE
        cell.fill = copy(base_fill) if base_fill is not None else PatternFill(fill_type=None)
    elif value == "A":
        cell.font = _FONT_AUSENTE
        cell.fill = _FILL_AUSENTE
    else:
        cell.font = _FONT_JUSTIFICADA
        cell.fill = _FILL_JUSTIFICADA
    return True


def _make_table(ws, name: str, last_col: int, last_row: int) -> None:
    table = Table(displayName=name, ref=f"A1:{get_column_letter(last_col)}{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,  # el color de cada fila lo da su sede
        showColumnStripes=False,
    )
    ws.add_table(table)


def _autosize(ws, columns: Sequence[str], last_row: int) -> None:
    for col_index, column_name in enumerate(columns, start=1):
        max_length = len(str(column_name))
        for row_idx in range(2, last_row + 1):
            value = ws.cell(row=row_idx, column=col_index).value
            if value is not None:
                max_length = max(max_length, len(str(value)))
        ws.column_dimensions[get_column_letter(col_index)].width = min(max_length + 4, 70)


def _write_table_sheet(wb: Workbook, title: str, columns: Sequence[str],
                       rows: List[Sequence[object]], empty_message: str,
                       table_name: str) -> None:
    """Agrega una pestaña con una tabla simple (encabezado + filas). Si no
    hay filas, deja una nota explicando que no hubo casos."""
    ws = wb.create_sheet(title[:31])
    ws.append(list(columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(list(row))

    last_row = ws.max_row
    if last_row == 1:
        ws.merge_cells(f"A2:{get_column_letter(len(columns))}2")
        nota = ws.cell(row=2, column=1)
        nota.value = empty_message
        nota.font = Font(italic=True, color="FF6A6A6A")
        last_row = 2
    else:
        _make_table(ws, table_name, len(columns), last_row)

    _autosize(ws, columns, last_row)
    ws.freeze_panes = "A2"


def _write_review_sheet(wb: Workbook, sheet_title: str, month_labels: List[str],
                        records_por_mes: Dict[str, List[Dict[str, object]]],
                        empty_message: str, table_name: str,
                        con_posibles: bool = False) -> None:
    """
    Pestaña de revisión manual para "No Encontrados" y "Ambiguos": una
    fila por cada participante de Zoom de esa categoría, con el mes, el
    nombre tal como lo escribió, la sede que escribió (si la hay), los
    minutos y —en "Ambiguos"— los socios del listado con los que podría
    corresponder.
    """
    columns = ["Mes", "Nombre en Zoom", "Sede en Zoom", "Duración (minutos)"]
    if con_posibles:
        columns.append("Posibles socios")
    rows = []
    for month_label in month_labels:
        for r in records_por_mes.get(month_label, []):
            row = [month_label, r.get("Nombre", ""), r.get("Sede", ""),
                   r.get("Duración (minutos)", 0)]
            if con_posibles:
                row.append("; ".join(r.get("Posibles", []) or []))
            rows.append(row)
    _write_table_sheet(wb, sheet_title, columns, rows, empty_message, table_name)


def _write_aproximados_sheet(wb: Workbook, month_labels: List[str],
                             aproximados_por_mes: Dict[str, List[Dict[str, object]]],
                             table_name: str = "TablaAproximados") -> None:
    """Pestaña "Cotejos aproximados": participantes que el programa SÍ
    asignó a un socio, pero no por nombre y apellido exactos (solo nombre
    único, apellido con otra ortografía, nombre escrito junto, desempate por
    sede...). Ya cuentan en la asistencia; se listan para poder revisarlos."""
    columns = ["Mes", "Nombre en Zoom", "Sede en Zoom", "Socio asignado",
               "Sede del socio", "Cómo se asignó", "Duración (minutos)"]
    rows = []
    for month_label in month_labels:
        for a in aproximados_por_mes.get(month_label, []):
            rows.append([
                month_label, a.get("Nombre", ""), a.get("Sede en Zoom", ""),
                a.get("Socio asignado", ""), a.get("Sede del socio", ""),
                a.get("Cómo se asignó", ""), a.get("Duración (minutos)", 0),
            ])
    _write_table_sheet(
        wb, SHEET_APROXIMADOS, columns, rows,
        "Todas las asignaciones fueron por nombre y apellido exactos.", table_name,
    )


def _write_pending_sheets(wb: Workbook, month_labels: List[str],
                          no_encontrados_por_mes, ambiguos_por_mes,
                          aproximados_por_mes) -> None:
    """Escribe las pestañas de revisión que correspondan (las que reciben
    datos, aunque vengan vacías)."""
    if ambiguos_por_mes is not None:
        _write_review_sheet(
            wb, SHEET_AMBIGUOS, month_labels, ambiguos_por_mes,
            empty_message="No hubo participantes ambiguos.",
            table_name="TablaAmbiguos", con_posibles=True,
        )
    if no_encontrados_por_mes is not None:
        _write_review_sheet(
            wb, SHEET_NO_ENCONTRADOS, month_labels, no_encontrados_por_mes,
            empty_message="No hubo participantes sin encontrar.",
            table_name="TablaNoEncontrados",
        )
    if aproximados_por_mes is not None:
        _write_aproximados_sheet(wb, month_labels, aproximados_por_mes)


def export_to_excel(records: List[Dict[str, object]], filepath: str,
                    sheet_name: str = "Asistencia", month_label: str = "",
                    no_encontrados: Optional[List[Dict[str, object]]] = None,
                    ambiguos: Optional[List[Dict[str, object]]] = None,
                    aproximados: Optional[List[Dict[str, object]]] = None) -> None:
    """
    Exporta el cotejo de UN mes a un libro de Excel:

      - Hoja "Asistencia": una fila por registro, en el orden recibido
        (para el cotejo, el mismo orden del listado de socios), con el
        color de cada sede y la asistencia P/A.
      - Si se reciben `no_encontrados`, `ambiguos` y/o `aproximados` (listas,
        aunque estén vacías), se agregan como pestañas del mismo libro:
        "Ambiguos", "No Encontrados" y "Cotejos aproximados".
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # Excel limita el nombre de hoja a 31 caracteres

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for record in records:
        try:
            duration_value = float(record.get("Duración (minutos)", 0))
        except (TypeError, ValueError):
            duration_value = 0.0
        ws.append([
            record.get("Nombre", ""),
            record.get("Apellido", ""),
            record.get("Sede", ""),
            duration_value,
            record.get("Asistencia", ""),
        ])

    last_row = ws.max_row
    duration_col = COLUMNS.index("Duración (minutos)") + 1
    attendance_col = COLUMNS.index("Asistencia") + 1

    if last_row >= 1:
        _make_table(ws, "TablaAsistencia", len(COLUMNS), last_row)

    fills = build_sede_fills(r.get("Sede", "") for r in records)
    for row_idx in range(2, last_row + 1):
        row_fill = fill_for(fills, ws.cell(row=row_idx, column=3).value or "")
        if row_fill is not None:
            for col in range(1, attendance_col):
                ws.cell(row=row_idx, column=col).fill = copy(row_fill)
        ws.cell(row=row_idx, column=duration_col).number_format = "0.00"
        apply_status_style(ws.cell(row=row_idx, column=attendance_col), row_fill)

    _autosize(ws, COLUMNS, last_row)
    ws.freeze_panes = "A2"

    label = month_label or ""
    _write_pending_sheets(
        wb, [label],
        None if no_encontrados is None else {label: no_encontrados},
        None if ambiguos is None else {label: ambiguos},
        None if aproximados is None else {label: aproximados},
    )
    wb.save(filepath)


def export_multi_month_to_excel(records: List[Dict[str, object]], month_labels: List[str],
                                filepath: str, sheet_name: str = "Asistencia",
                                no_encontrados_por_mes: Optional[Dict[str, List[Dict[str, object]]]] = None,
                                ambiguos_por_mes: Optional[Dict[str, List[Dict[str, object]]]] = None,
                                aproximados_por_mes: Optional[Dict[str, List[Dict[str, object]]]] = None) -> None:
    """
    Igual que export_to_excel, pero para el "Reporte Multi-Mes": en vez de
    una sola columna "Asistencia", UNA COLUMNA POR MES (mismo formato de
    colores), en el orden dado por `month_labels`, con una fila por socio
    en el mismo orden del listado y el color de su sede.

    `records` trae "Nombre", "Apellido", "Sede" y una llave por cada mes
    con "P" o "A" (formato de multi_month_processor.build_multi_month_matrix).

    Las pestañas "Ambiguos", "No Encontrados" y "Cotejos aproximados" se
    agregan al mismo libro cuando se reciben sus datos (aunque vengan
    vacíos), mes por mes.
    """
    columns = ["Nombre", "Apellido", "Sede"] + list(month_labels)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for record in records:
        row = [record.get("Nombre", ""), record.get("Apellido", ""), record.get("Sede", "")]
        row += [record.get(label, "") for label in month_labels]
        ws.append(row)

    last_row = ws.max_row
    if last_row >= 1:
        _make_table(ws, "TablaAsistenciaMultiMes", len(columns), last_row)

    first_month_col = 4  # A=Nombre, B=Apellido, C=Sede -> D en adelante = meses
    fills = build_sede_fills(r.get("Sede", "") for r in records)
    for row_idx in range(2, last_row + 1):
        row_fill = fill_for(fills, ws.cell(row=row_idx, column=3).value or "")
        if row_fill is not None:
            for col in range(1, first_month_col):
                ws.cell(row=row_idx, column=col).fill = copy(row_fill)
        for offset in range(len(month_labels)):
            cell = ws.cell(row=row_idx, column=first_month_col + offset)
            cell.alignment = Alignment(horizontal="center")
            apply_status_style(cell, row_fill)

    _autosize(ws, columns, last_row)
    ws.freeze_panes = "D2"

    _write_pending_sheets(
        wb, list(month_labels), no_encontrados_por_mes, ambiguos_por_mes, aproximados_por_mes,
    )
    wb.save(filepath)


def export_pendientes_to_excel(filepath: str, month_label: str,
                               no_encontrados: List[Dict[str, object]],
                               ambiguos: List[Dict[str, object]],
                               aproximados: Optional[List[Dict[str, object]]] = None) -> None:
    """Libro de Excel solo con las pestañas para revisar a mano de UN mes
    ("Ambiguos", "No Encontrados" y, si se entrega, "Cotejos aproximados")."""
    wb = Workbook()
    wb.remove(wb.active)
    label = month_label or ""
    _write_pending_sheets(
        wb, [label], {label: no_encontrados}, {label: ambiguos},
        None if aproximados is None else {label: aproximados},
    )
    wb.save(filepath)
